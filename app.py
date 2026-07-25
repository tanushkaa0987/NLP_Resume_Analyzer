import os
import re
import tempfile
from collections import Counter
from pathlib import Path

import pdfplumber
import PyPDF2
from docx import Document
from flask import Flask, jsonify, render_template, request

try:
    import spacy
except ImportError:
    spacy = None

try:
    from sklearn.feature_extraction.text import TfidfVectorizer
except ImportError:
    TfidfVectorizer = None


BASE_DIR = Path(__file__).resolve().parent
UPLOAD_EXTENSIONS = {".pdf", ".docx"}
MAX_UPLOAD_MB = 8

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_MB * 1024 * 1024


JOB_KEYWORDS = {
    "software_engineer": [
        "python",
        "javascript",
        "java",
        "react",
        "node",
        "api",
        "sql",
        "git",
        "docker",
        "aws",
        "testing",
        "microservices",
        "flask",
        "django",
    ],
    "data_scientist": [
        "python",
        "machine learning",
        "statistics",
        "pandas",
        "numpy",
        "scikit learn",
        "tensorflow",
        "pytorch",
        "sql",
        "visualization",
        "nlp",
        "modeling",
    ],
    "product_manager": [
        "roadmap",
        "stakeholder",
        "analytics",
        "strategy",
        "user research",
        "prioritization",
        "kpi",
        "launch",
        "experimentation",
        "agile",
    ],
    "default": [
        "leadership",
        "communication",
        "project",
        "analysis",
        "collaboration",
        "problem solving",
        "strategy",
        "planning",
        "delivery",
        "metrics",
    ],
}

ACTION_VERBS = {
    "built",
    "created",
    "designed",
    "developed",
    "delivered",
    "improved",
    "reduced",
    "increased",
    "launched",
    "managed",
    "led",
    "automated",
    "optimized",
    "implemented",
    "analyzed",
}


def load_nlp():
    if spacy is None:
        return None

    try:
        return spacy.load("en_core_web_sm")
    except OSError:
        return spacy.blank("en")


NLP = load_nlp()


def normalize_text(text):
    text = re.sub(r"\s+", " ", text or "")
    return text.strip()


def extract_pdf_text(path):
    pages = []

    try:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
    except Exception:
        pages = []

    text = "\n".join(pages).strip()
    if text:
        return text

    with open(path, "rb") as resume_file:
        reader = PyPDF2.PdfReader(resume_file)
        return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_docx_text(path):
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_text = []

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                table_text.append(cell.text)

    return "\n".join(paragraphs + table_text)


def extract_resume_text(uploaded_file):
    suffix = Path(uploaded_file.filename).suffix.lower()
    if suffix not in UPLOAD_EXTENSIONS:
        raise ValueError("Please upload a PDF or DOCX resume.")

    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        uploaded_file.save(tmp.name)
        tmp_path = tmp.name

    try:
        if suffix == ".pdf":
            return normalize_text(extract_pdf_text(tmp_path))
        return normalize_text(extract_docx_text(tmp_path))
    finally:
        os.remove(tmp_path)


def infer_target_role(text):
    text_lower = text.lower()
    scores = {}
    for role, keywords in JOB_KEYWORDS.items():
        if role == "default":
            continue
        scores[role] = sum(1 for keyword in keywords if keyword in text_lower)

    best_role = max(scores, key=scores.get)
    return best_role if scores[best_role] > 1 else "default"


def keyword_match_score(text, role):
    text_lower = text.lower()
    keywords = JOB_KEYWORDS.get(role, JOB_KEYWORDS["default"])
    matched = [keyword for keyword in keywords if keyword in text_lower]
    missing = [keyword for keyword in keywords if keyword not in text_lower]
    score = round((len(matched) / len(keywords)) * 100)
    return score, matched, missing


def readability_score(text):
    words = re.findall(r"[A-Za-z]+", text)
    sentences = re.split(r"[.!?]+", text)
    sentences = [sentence for sentence in sentences if sentence.strip()]

    if not words or not sentences:
        return 0

    avg_sentence_length = len(words) / len(sentences)
    score = 100 - max(0, avg_sentence_length - 18) * 3
    return max(40, min(100, round(score)))


def impact_score(text):
    bullets = re.findall(r"(?:^|\n)\s*(?:[-*•]|\d+[.)])\s+", text)
    numbers = re.findall(r"\b(?:\d+[%+$]?|\$[\d,.]+|[0-9]+x)\b", text)
    lower_text = text.lower()
    verbs_found = [verb for verb in ACTION_VERBS if re.search(rf"\b{verb}\b", lower_text)]

    score = 45 + min(25, len(numbers) * 4) + min(20, len(verbs_found) * 2) + min(10, len(bullets))
    return min(100, score), numbers[:8], sorted(verbs_found)[:10]


def extract_skills(text):
    vocabulary = sorted(set(sum(JOB_KEYWORDS.values(), [])))
    if TfidfVectorizer is None:
        text_lower = text.lower()
        return [keyword for keyword in vocabulary if keyword in text_lower][:12]

    vectorizer = TfidfVectorizer(vocabulary=vocabulary, ngram_range=(1, 2), stop_words="english")
    matrix = vectorizer.fit_transform([text.lower()])
    feature_names = vectorizer.get_feature_names_out()
    weights = matrix.toarray()[0]
    ranked = sorted(zip(feature_names, weights), key=lambda item: item[1], reverse=True)
    return [name for name, weight in ranked if weight > 0][:12]


def extract_entities(text):
    if NLP is None:
        return {"people": [], "organizations": [], "locations": []}

    doc = NLP(text[:20000])
    entities = {"people": [], "organizations": [], "locations": []}

    for entity in doc.ents:
        value = entity.text.strip()
        if len(value) < 2:
            continue
        if entity.label_ == "PERSON":
            entities["people"].append(value)
        elif entity.label_ == "ORG":
            entities["organizations"].append(value)
        elif entity.label_ in {"GPE", "LOC"}:
            entities["locations"].append(value)

    return {key: list(dict.fromkeys(values))[:5] for key, values in entities.items()}


def build_recommendations(text, missing_keywords, matched_keywords, impact_numbers):
    recommendations = []

    if len(text.split()) < 250:
        recommendations.append("Add more detail to your recent roles, projects, tools, and measurable outcomes.")

    if missing_keywords:
        recommendations.append(
            "Add role-relevant keywords where truthful: " + ", ".join(missing_keywords[:6]) + "."
        )

    if len(impact_numbers) < 3:
        recommendations.append("Quantify achievements with metrics such as percentage gains, time saved, revenue, users, or cost reduction.")

    if len(matched_keywords) < 5:
        recommendations.append("Include a focused skills section near the top so recruiters and ATS systems can scan your fit quickly.")

    if not re.search(r"\b(?:email|phone|linkedin|github|portfolio)\b", text.lower()):
        recommendations.append("Make sure contact details and professional links are visible in the header.")

    if not recommendations:
        recommendations.append("Strong foundation. Tighten the top summary and tailor keywords to each job description before applying.")

    return recommendations[:5]


def analyze_resume(text):
    if len(text.split()) < 40:
        raise ValueError("The resume text is too short to analyze. Try another file with selectable text.")

    target_role = infer_target_role(text)
    keyword_score, matched_keywords, missing_keywords = keyword_match_score(text, target_role)
    clarity = readability_score(text)
    impact, impact_numbers, action_verbs = impact_score(text)
    skills = extract_skills(text)
    entities = extract_entities(text)

    section_names = ["experience", "education", "skills", "projects", "certifications", "summary"]
    found_sections = [section.title() for section in section_names if re.search(rf"\b{section}\b", text, re.I)]
    section_score = round((len(found_sections) / len(section_names)) * 100)

    overall = round(keyword_score * 0.35 + clarity * 0.2 + impact * 0.3 + section_score * 0.15)
    word_count = len(re.findall(r"\b\w+\b", text))
    common_terms = Counter(re.findall(r"\b[a-zA-Z]{4,}\b", text.lower())).most_common(8)

    return {
        "overall_score": overall,
        "target_role": target_role.replace("_", " ").title(),
        "scores": {
            "keyword_match": keyword_score,
            "clarity": clarity,
            "impact": impact,
            "resume_structure": section_score,
        },
        "matched_keywords": matched_keywords,
        "missing_keywords": missing_keywords[:10],
        "detected_skills": skills,
        "found_sections": found_sections,
        "impact_numbers": impact_numbers,
        "action_verbs": action_verbs,
        "entities": entities,
        "word_count": word_count,
        "common_terms": [{"term": term, "count": count} for term, count in common_terms],
        "recommendations": build_recommendations(text, missing_keywords, matched_keywords, impact_numbers),
        "preview": text[:700] + ("..." if len(text) > 700 else ""),
    }


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/analyze", methods=["POST"])
def analyze():
    if "resume" not in request.files:
        return jsonify({"error": "No resume file was uploaded."}), 400

    uploaded_file = request.files["resume"]
    if not uploaded_file.filename:
        return jsonify({"error": "Choose a resume file first."}), 400

    try:
        text = extract_resume_text(uploaded_file)
        result = analyze_resume(text)
        return jsonify(result)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        return jsonify({"error": f"Could not analyze this resume: {exc}"}), 500


if __name__ == "__main__":
    app.run(debug=True)
